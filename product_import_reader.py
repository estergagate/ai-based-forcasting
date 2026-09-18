"""Read a bounded Excel sheet and validate product rows. No database writes."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook


def read_excel(data):
    if len(data) > 5 * 1024 * 1024:
        raise ValueError('Use a file smaller than 5 MB.')
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 25 * 1024 * 1024:
                raise ValueError('This workbook is too large. Use a smaller copy.')
    except BadZipFile:
        raise ValueError('Choose a valid Excel .xlsx file.')
    book = None
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        sheet = book.worksheets[0]
        # Ignore incorrect saved worksheet dimensions; bound actual iteration instead.
        sheet.reset_dimensions()
        rows = []
        for row in sheet.iter_rows(max_row=522, max_col=51):
            if len(rows) >= 521:
                if any(cell.value is not None for cell in row):
                    raise ValueError('Use at most 500 product rows after the headings.')
                break
            if row[50].value is not None:
                raise ValueError('Use at most 50 columns in the first worksheet.')
            result = []
            for cell in row[:50]:
                value = '' if cell.value is None else str(cell.value).strip()
                if len(value) > 1000:
                    raise ValueError('A cell contains too much text. Use a simple product list.')
                result.append({'value': value, 'invalid': cell.data_type in ('f', 'e')})
            rows.append(result)
        while rows and not any(cell['value'] for cell in rows[-1]):
            rows.pop()
        if not rows:
            raise ValueError('The first worksheet is empty.')
        width = max(i + 1 for row in rows for i, cell in enumerate(row) if cell['value'])
        return {'sheet': sheet.title, 'rows': [row[:width] for row in rows]}
    except ValueError:
        raise
    except Exception:
        raise ValueError('This workbook could not be read. Save a fresh .xlsx copy and try again.')
    finally:
        if book is not None:
            book.close()


def validate_products(rows, units):
    lookup = {unit.casefold(): unit for unit in units}
    aliases = {'pcs': 'Piece', 'pc': 'Piece', 'pieces': 'Piece',
               'servings': 'Serving', 'bottles': 'Bottle'}
    errors, clean, seen = [], [], set()
    for number, row in enumerate(rows, 1):
        label = f"Row {row.get('source', number)}"
        name = str(row.get('name', '')).strip()
        unit = str(row.get('unit', '')).strip()
        price_text = str(row.get('price', '')).strip()
        if not 1 <= len(name) <= 100:
            errors.append(f'{label}: enter a product name of 1–100 characters.')
        if name.casefold() in seen:
            errors.append(f'{label}: {name} appears more than once. Keep one row per product.')
        seen.add(name.casefold())
        unit = lookup.get(unit.casefold(), lookup.get(aliases.get(unit.casefold(), '').casefold()))
        if unit is None:
            errors.append(f'{label}: choose a unit from your saved units.')
        price = None
        try:
            if len(price_text) > 30:
                raise ValueError
            price = Decimal(price_text)
            if (not price.is_finite() or price < 0 or price > Decimal('99999999.99')
                    or price != price.quantize(Decimal('0.01'))):
                raise ValueError
        except (ValueError, InvalidOperation):
            errors.append(f'{label}: enter a peso price from 0 to 99,999,999.99, with up to two decimal places. No currency symbols or commas.')
        clean.append((name, unit, price))
    if not rows:
        errors.append('Add at least one product.')
    if len(rows) > 500:
        errors.append('Import at most 500 products at a time.')
    return clean, errors
