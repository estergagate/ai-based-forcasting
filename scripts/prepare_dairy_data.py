import calendar
import csv
import math
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook


# Find the project folder regardless of where the command is run.
BASE_DIR = Path(__file__).resolve().parents[1]

INPUT_DIR = BASE_DIR / "data" / "raw" / "dairy" / "product 1"
OUTPUT_DIR = BASE_DIR / "data" / "processed"

MONTHS = {
    name.lower(): number
    for number, name in enumerate(calendar.month_name)
    if name
}


def prepare_data():
    records = {}

    for year in (2020, 2021, 2022):
        file_path = INPUT_DIR / f"product 1 {year}.xlsx"

        if not file_path.exists():
            raise FileNotFoundError(f"Missing file: {file_path}")

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows)

            columns = {
                str(name).strip(): index
                for index, name in enumerate(headers)
                if name is not None
            }

            required = ("Day", "Month", "Year", "daily_unit_sales")

            for name in required:
                if name not in columns:
                    raise ValueError(
                        f"{file_path.name}: missing column {name}"
                    )

            for row_number, row in enumerate(rows, start=2):
                if all(value is None for value in row):
                    continue

                sales_date = date(
                    int(row[columns["Year"]]),
                    MONTHS[str(row[columns["Month"]]).strip().lower()],
                    int(row[columns["Day"]]),
                )

                if sales_date.year != year:
                    raise ValueError(
                        f"Unexpected year in {file_path.name}, "
                        f"row {row_number}"
                    )

                if sales_date in records:
                    raise ValueError(f"Duplicate date: {sales_date}")

                raw_quantity = row[columns["daily_unit_sales"]]

                try:
                    quantity = float(raw_quantity)

                    if not math.isfinite(quantity):
                        raise ValueError

                    status = (
                        "negative_review" if quantity < 0 else "recorded"
                    )

                except (TypeError, ValueError):
                    quantity = ""
                    status = "invalid_quantity"

                records[sales_date] = {
                    "quantity": quantity,
                    "status": status,
                    "source_file": file_path.name,
                    "source_row": row_number,
                }

        finally:
            workbook.close()

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / "dairy_product_1_review.csv"

    counts = {}
    current_date = date(2020, 1, 1)
    end_date = date(2022, 12, 31)

    with output_path.open("w", newline="", encoding="utf-8") as output:
        writer = csv.writer(output)
        writer.writerow([
            "date",
            "product",
            "quantity",
            "status",
            "source_file",
            "source_row",
        ])

        while current_date <= end_date:
            record = records.get(current_date)

            if record is None:
                record = {
                    "quantity": "",
                    "status": "missing_date",
                    "source_file": "",
                    "source_row": "",
                }

            writer.writerow([
                current_date.isoformat(),
                "dairy_product_1",
                record["quantity"],
                record["status"],
                record["source_file"],
                record["source_row"],
            ])

            status = record["status"]
            counts[status] = counts.get(status, 0) + 1
            current_date += timedelta(days=1)

    print(f"Created: {output_path}")
    print(f"Source records: {len(records)}")

    for status, count in sorted(counts.items()):
        print(f"{status}: {count}")

    print("Review file prepared. No model has been trained yet.")


if __name__ == "__main__":
    prepare_data()